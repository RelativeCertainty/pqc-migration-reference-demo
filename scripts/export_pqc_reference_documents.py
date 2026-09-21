#!/usr/bin/env python3
"""Create editable, complete Word projections of local PQC review documents.

Offline operator utility: no source-system calls, dependency acquisition, upload,
or authority to accept an assessment. The immutable Markdown remains the source.
Use the document skill's render-and-inspect gate after running this exporter.

Layout: standard_business_brief, memo_masthead, with named detail_appendix and
literal_code overrides. All headings, list markers and hyperlinks are native
Word objects; there are no tables-as-layout, decorative borders or fake bullets.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from lxml import etree


STYLE_TOKENS = {
    "preset": "standard_business_brief",
    "header_pattern": "memo_masthead",
    "page": {"width_dxa": 12240, "height_dxa": 15840, "margin_dxa": 1440},
    "header_footer_inches": 0.492,
    "body": {"font": "Calibri", "pt": 11, "before": 0, "after": 6, "line": 264},
    "heading_1": {"pt": 16, "color": "2E74B5", "before": 16, "after": 8},
    "heading_2": {"pt": 13, "color": "2E74B5", "before": 12, "after": 6},
    "heading_3": {"pt": 12, "color": "1F4D78", "before": 8, "after": 4},
    "list": {"left_dxa": 720, "hanging_dxa": 360, "after": 8, "line": 280},
    "named_overrides": {
        "memo_title": {"pt": 23, "color": "000000", "before": 0, "after": 4},
        "memo_subtitle": {"pt": 11, "color": "555555", "before": 0, "after": 8},
        "memo_metadata": {"pt": 10, "before": 0, "after": 2, "line": 264},
        "quiet_furniture": {"pt": 9, "color": "555555", "after": 0, "line": 240},
        "detail_appendix": {"pt": 10, "before": 0, "after": 2, "line": 240},
        "detail_record_heading": {
            "pt": 11,
            "before": 8,
            "after": 3,
            "line": 240,
            "color": "1F4D78",
        },
        "literal_code": {"font": "Consolas", "pt": 9.5, "after": 0, "line": 240},
        "inline_code": {"font": "Consolas", "pt": 9.5},
        "no_decorative_rules": True,
    },
}

DETAIL_SECTIONS = {
    "Inventory and relationship appendix",
    "Cryptographic-use review register",
}
HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
LIST = re.compile(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$")
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!!)\[[^\]]+\]\((?:<[^>]+>|[^)])+\))")
LINK = re.compile(r"^\[([^\]]+)\]\((.*)\)$")


@dataclass(frozen=True)
class Block:
    kind: str
    text: str
    level: int = 0
    start: int = 1


def parse_markdown(text: str) -> list[Block]:
    """Parse the intentionally small, lossless subset used in PQC reports.

    Unsupported table/image/HTML syntax fails rather than silently losing it.
    Markdown soft line breaks become spaces; fenced text keeps literal lines.
    """
    blocks: list[Block] = []
    paragraph: list[str] = []
    fence: str | None = None

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph)))
            paragraph.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if fence:
            if stripped.startswith(fence):
                fence = None
            else:
                blocks.append(Block("code", line))
            continue
        if stripped.startswith(("```", "~~~")):
            flush()
            fence = stripped[:3]
            continue
        if not stripped:
            flush()
            continue
        if stripped.startswith(("|", "![", "<table", "<img")):
            raise ValueError("unsupported Markdown table/image/HTML: adapt explicitly")
        if set(stripped) <= {"-", "*", "_"} and len(stripped) >= 3:
            flush()
            continue  # A Markdown thematic break is whitespace, not a border.
        match = HEADING.match(line)
        if match:
            flush()
            blocks.append(Block("heading", match[2], len(match[1])))
            continue
        match = LIST.match(line)
        if match:
            flush()
            numeric = match[2][0].isdigit()
            blocks.append(
                Block(
                    "number" if numeric else "bullet",
                    match[3],
                    min(len(match[1].expandtabs(4)) // 2, 8),
                    int(match[2][:-1]) if numeric else 1,
                )
            )
            continue
        if stripped.startswith("> "):
            flush()
            blocks.append(Block("quote", stripped[2:]))
            continue
        if line.startswith("  ") and blocks and blocks[-1].kind in {"number", "bullet"}:
            prev = blocks[-1]
            blocks[-1] = Block(
                prev.kind, prev.text + "\n" + stripped, prev.level, prev.start
            )
            continue
        paragraph.append(stripped)
    flush()
    if fence:
        raise ValueError("unclosed Markdown code fence")
    if not blocks or blocks[0].kind != "heading" or blocks[0].level != 1:
        raise ValueError("document must begin with exactly one level-one title")
    if any(block.kind == "heading" and block.level == 1 for block in blocks[1:]):
        raise ValueError("additional level-one headings require explicit conversion")
    return blocks


def _child(parent: Any, tag: str, attrs: dict[str, Any] | None = None) -> Any:
    node = OxmlElement(tag)
    for name, value in (attrs or {}).items():
        node.set(qn(name), str(value))
    parent.append(node)
    return node


def _style(
    document: Any,
    name: str,
    *,
    pt: float = 11,
    color: str = "000000",
    before: float = 0,
    after: float = 6,
    line: int = 264,
    font: str = "Calibri",
    bold: bool = False,
    heading: bool = False,
) -> Any:
    styles = document.styles
    style = (
        styles[name]
        if name in styles
        else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    style.base_style = None if name == "Normal" else styles["Normal"]
    style.font.name = font
    style.font.size = Pt(pt)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = False
    style.font.underline = False
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + attr), font)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line / 240
    fmt.left_indent = Twips(0)
    fmt.right_indent = Twips(0)
    fmt.first_line_indent = Twips(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.keep_with_next = heading
    fmt.keep_together = False
    fmt.widow_control = True
    fmt.page_break_before = False
    for border in list(style.element.findall(qn("w:pPr") + "/" + qn("w:pBdr"))):
        border.getparent().remove(border)
    return style


def _configure(document: Any) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    _style(document, "Normal")
    _style(document, "Title", pt=23, before=0, after=4, bold=True, heading=True)
    _style(document, "Subtitle", color="555555", after=8, heading=True)
    for level in (1, 2, 3):
        _style(
            document,
            f"Heading {level}",
            **STYLE_TOKENS[f"heading_{level}"],
            bold=True,
            heading=True,
        )
    _style(document, "PQC Metadata", pt=10, after=2)
    _style(document, "PQC Furniture", pt=9, color="555555", after=0, line=240)
    _style(document, "PQC List", after=8, line=280)
    _style(document, "PQC Detail", pt=10, after=2, line=240)
    _style(document, "PQC Detail List", pt=10, after=2, line=240)
    detail = _style(
        document,
        "PQC Detail Heading",
        pt=11,
        before=8,
        after=3,
        line=240,
        color="1F4D78",
        bold=True,
        heading=True,
    )
    _child(detail.element.get_or_add_pPr(), "w:outlineLvl", {"w:val": 1})
    _style(document, "PQC Code", pt=9.5, font="Consolas", after=0, line=240)
    _style(document, "PQC Quote", color="1F4D78", after=6)
    _style(document, "PQC Table Citation", before=4, after=4)
    for name in ("PQC List", "PQC Detail List"):
        document.styles[name].paragraph_format.left_indent = Twips(720)
        document.styles[name].paragraph_format.first_line_indent = Twips(-360)
    header = section.header.paragraphs[0]
    header.style = document.styles["PQC Furniture"]
    header.text = "PQC ENTERPRISE REFERENCE  |  SYNTHETIC DEVELOPMENT EVIDENCE"
    footer = section.footer.paragraphs[0]
    footer.style = document.styles["PQC Furniture"]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Draft · No enterprise acceptance or execution authority  |  ")
    for field, suffix in (("PAGE", " of "), ("NUMPAGES", "")):
        element = _child(footer._p, "w:fldSimple", {"w:instr": field})
        _child(_child(element, "w:r"), "w:t").text = "1"
        footer.add_run(suffix)
    document.core_properties.author = "PQC reference engineering"
    document.core_properties.last_modified_by = "PQC reference document exporter"
    document.core_properties.subject = (
        "Synthetic development evidence; not enterprise acceptance"
    )
    document.core_properties.keywords = "synthetic, development, evidence, draft"
    # Remove inactive stock-template rules too: editing with a quotation style
    # later must not reintroduce the decorative borders rejected for this pack.
    for border in list(document.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)


class _Numbering:
    def __init__(self, document: Any) -> None:
        self.root = document.part.numbering_part.element
        self.abstract_id = (
            max(
                (
                    int(n.get(qn("w:abstractNumId")))
                    for n in self.root.findall(qn("w:abstractNum"))
                ),
                default=-1,
            )
            + 1
        )
        self.num_id = (
            max(
                (int(n.get(qn("w:numId"))) for n in self.root.findall(qn("w:num"))),
                default=0,
            )
            + 1
        )

    def new(self, kind: str, start: int = 1, detail: bool = False) -> int:
        abstract_id, num_id = self.abstract_id, self.num_id
        self.abstract_id += 1
        self.num_id += 1
        abstract = _child(self.root, "w:abstractNum", {"w:abstractNumId": abstract_id})
        first_number = self.root.find(qn("w:num"))
        if first_number is not None:
            self.root.remove(abstract)
            self.root.insert(self.root.index(first_number), abstract)
        _child(abstract, "w:multiLevelType", {"w:val": "multilevel"})
        for level in range(9):
            lvl = _child(abstract, "w:lvl", {"w:ilvl": level})
            _child(lvl, "w:start", {"w:val": start if level == 0 else 1})
            _child(
                lvl, "w:numFmt", {"w:val": "decimal" if kind == "number" else "bullet"}
            )
            _child(
                lvl,
                "w:lvlText",
                {"w:val": f"%{level + 1}." if kind == "number" else "•"},
            )
            _child(lvl, "w:lvlJc", {"w:val": "left"})
            ppr = _child(lvl, "w:pPr")
            left = 720 + level * 360
            _child(_child(ppr, "w:tabs"), "w:tab", {"w:val": "num", "w:pos": left})
            _child(
                ppr,
                "w:spacing",
                {
                    "w:before": 0,
                    "w:after": 40 if detail else 160,
                    "w:line": 240 if detail else 280,
                    "w:lineRule": "auto",
                },
            )
            _child(ppr, "w:ind", {"w:left": left, "w:hanging": 360})
            rpr = _child(lvl, "w:rPr")
            _child(rpr, "w:rFonts", {"w:ascii": "Calibri", "w:hAnsi": "Calibri"})
            _child(rpr, "w:sz", {"w:val": 20 if detail else 22})
        number = _child(self.root, "w:num", {"w:numId": num_id})
        _child(number, "w:abstractNumId", {"w:val": abstract_id})
        return num_id


def _inline(paragraph: Any, text: str) -> None:
    def run(value: str, *, bold: bool = False, code: bool = False) -> None:
        node = paragraph.add_run(value)
        node.bold = bold
        if code:
            node.font.name = "Consolas"
            node.font.size = Pt(9.5)
            fonts = node._r.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:ascii"), "Consolas")
            fonts.set(qn("w:hAnsi"), "Consolas")

    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run(token[1:-1], code=True)
        elif token.startswith("**") and token.endswith("**"):
            run(token[2:-2], bold=True)
        elif match := LINK.match(token):
            label, target = match.groups()
            target = (
                target[1:-1]
                if target.startswith("<") and target.endswith(">")
                else target
            )
            if urlsplit(target).scheme not in {"", "https", "http", "mailto"}:
                raise ValueError("unsupported hyperlink scheme")
            from docx.opc.constants import RELATIONSHIP_TYPE

            relationship = paragraph.part.relate_to(
                target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
            hyperlink = _child(paragraph._p, "w:hyperlink", {"r:id": relationship})
            rpr = _child(_child(hyperlink, "w:r"), "w:rPr")
            _child(rpr, "w:color", {"w:val": "1F4D78"})
            _child(rpr, "w:u", {"w:val": "single"})
            _child(rpr.getparent(), "w:t").text = label
        else:
            run(token)


def _plain_inline(text: str) -> str:
    for token in INLINE.findall(text):
        if token.startswith("`"):
            replacement = token[1:-1]
        elif token.startswith("**"):
            replacement = token[2:-2]
        else:
            match = LINK.match(token)
            assert match is not None
            replacement = match[1]
        text = text.replace(token, replacement, 1)
    return text


def _audit_preset(document: Any) -> dict[str, Any]:
    section = document.sections[0]
    geometry = (
        section.page_width.twips,
        section.page_height.twips,
        section.top_margin.twips,
        section.bottom_margin.twips,
        section.left_margin.twips,
        section.right_margin.twips,
    )
    if geometry != (12240, 15840, 1440, 1440, 1440, 1440):
        raise ValueError("document page geometry differs from preset")
    for name, tokens in [("Normal", STYLE_TOKENS["body"])] + [
        (f"Heading {i}", STYLE_TOKENS[f"heading_{i}"]) for i in (1, 2, 3)
    ]:
        style = document.styles[name]
        fmt = style.paragraph_format
        if (
            style.font.name != "Calibri"
            or style.font.size.pt != tokens["pt"]
            or fmt.space_before.pt != tokens["before"]
            or fmt.space_after.pt != tokens["after"]
            or round(fmt.line_spacing * 240) != tokens.get("line", 264)
        ):
            raise ValueError("document paragraph styles differ from preset")
        if "color" in tokens and str(style.font.color.rgb) != tokens["color"]:
            raise ValueError("document heading colors differ from preset")
    lists = [
        p
        for p in document.paragraphs
        if p.style.name in {"PQC List", "PQC Detail List"}
    ]
    if any(p._p.find(qn("w:pPr") + "/" + qn("w:numPr")) is None for p in lists):
        raise ValueError("list paragraph missing native numbering")
    seen_number = False
    for node in document.part.numbering_part.element:
        if node.tag == qn("w:num"):
            seen_number = True
        elif node.tag == qn("w:abstractNum") and seen_number:
            raise ValueError("numbering part violates abstract-before-concrete ordering")
    return {
        "page_geometry": "pass",
        "styles": "pass",
        "native_lists": "pass",
        "decorative_borders": "none",
        "named_overrides": sorted(STYLE_TOKENS["named_overrides"]),
    }


def export_document(
    source: Path, output: Path, *, replace_existing: bool = False
) -> dict[str, Any]:
    source, output = source.resolve(), output.resolve()
    if source.suffix != ".md" or output.suffix != ".docx" or source == output:
        raise ValueError("expected a Markdown source and a separate DOCX output")
    if output.exists() and not replace_existing:
        raise FileExistsError("refusing to overwrite existing document")
    source_bytes = source.read_bytes()
    blocks = parse_markdown(source_bytes.decode("utf-8"))
    document = Document()
    _configure(document)
    document.core_properties.title = _plain_inline(blocks[0].text)
    _inline(document.add_paragraph(style="Title"), blocks[0].text)
    document.add_paragraph(
        "Editable review copy · Complete source content retained", "Subtitle"
    )
    metadata = document.add_paragraph(style="PQC Metadata")
    metadata.add_run("Source: ").bold = True
    _inline(metadata, f"[{source.name}]({source.name})")
    digest = hashlib.sha256(source_bytes).hexdigest()
    document.add_paragraph(f"Source SHA-256: {digest}", "PQC Metadata")
    if source.stem in {"phase1-report", "phase2-report"}:
        p = document.add_paragraph(style="PQC Metadata")
        _inline(
            p,
            "Companion: [full machine-readable report pack](report-pack.json). "
            "Detailed records follow the narrative in this editable document.",
        )
    numbering = _Numbering(document)
    previous_kind: str | None = None
    current_num = 0
    detail = False
    expected: list[str] = [_plain_inline(blocks[0].text)]
    for block in blocks[1:]:
        if block.kind == "heading":
            previous_kind = None
            if block.level == 2:
                detail = block.text in DETAIL_SECTIONS
            level = min(block.level - 1, 3)
            style = (
                "PQC Detail Heading"
                if detail and block.level > 2
                else f"Heading {level}"
            )
            paragraph = document.add_paragraph(style=style)
            if block.level == 2 and block.text in DETAIL_SECTIONS:
                paragraph.paragraph_format.page_break_before = True
        elif block.kind in {"number", "bullet"}:
            if previous_kind != block.kind:
                current_num = numbering.new(block.kind, block.start, detail)
            paragraph = document.add_paragraph(
                style="PQC Detail List" if detail else "PQC List"
            )
            numpr = _child(paragraph._p.get_or_add_pPr(), "w:numPr")
            _child(numpr, "w:ilvl", {"w:val": block.level})
            _child(numpr, "w:numId", {"w:val": current_num})
            if block.level:
                paragraph.paragraph_format.left_indent = Twips(720 + block.level * 360)
            previous_kind = block.kind
        else:
            previous_kind = None
            style = {"code": "PQC Code", "quote": "PQC Quote"}.get(
                block.kind, "PQC Detail" if detail else "Normal"
            )
            paragraph = document.add_paragraph(style=style)
        if block.kind == "code":
            paragraph.add_run(block.text)
            expected.append(block.text)
        else:
            _inline(paragraph, block.text)
            expected.append(_plain_inline(block.text))
    output.parent.mkdir(parents=True, exist_ok=True)
    # Word's optional effects stylesheet must not retain a stale Title or
    # quotation design. This deck-free document uses the same native styles in
    # both Office representations, without effect-only formatting.
    for part in document.part.package.parts:
        if str(part.partname) == "/word/stylesWithEffects.xml":
            part._blob = etree.tostring(
                document.styles.element,
                encoding="UTF-8",
                xml_declaration=True,
                standalone=True,
            )
    document.save(output)
    # Structural readback is not visual QA. It prevents lost text and bad packages.
    reread = Document(output)
    preset_audit = _audit_preset(reread)
    visible = [p.text for p in reread.paragraphs]
    cursor = 0
    for value in expected:
        while cursor < len(visible) and visible[cursor] != value:
            cursor += 1
        if cursor == len(visible):
            raise ValueError("export content readback mismatch")
        cursor += 1
    with ZipFile(output) as archive:
        if archive.testzip() is not None:
            raise ValueError("invalid DOCX package")
        for part in archive.namelist():
            if part.endswith(".xml") and b"<w:pBdr" in archive.read(part):
                raise ValueError("unexpected decorative paragraph border")
    return {
        "source": str(source),
        "source_sha256": digest,
        "output": str(output),
        "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "source_blocks": len(blocks),
        "content_readback": "pass",
        "native_lists": sum(b.kind in {"bullet", "number"} for b in blocks),
        "preset": STYLE_TOKENS,
        "preset_audit": preset_audit,
        "visual_review": "required_not_yet_performed",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("sources", type=Path, nargs="+")
    parser.add_argument(
        "--replace-generated",
        action="store_true",
        help="explicitly replace DOCX files previously generated from these sources",
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        required=True,
        help="new JSON structural evidence manifest (not a visual-QA assertion)",
    )
    args = parser.parse_args()
    if args.manifest.exists() and not args.replace_generated:
        parser.error("manifest already exists")
    exports = [
        export_document(
            p, p.with_suffix(".docx"), replace_existing=args.replace_generated
        )
        for p in args.sources
    ]
    manifest = {
        "schema_version": "pqc.document-export.v1",
        "exporter_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "exports": exports,
        "network_operations": False,
        "acceptance_authority": False,
    }
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    args.manifest.write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "documents": len(exports),
                "content_readback": "pass",
                "visual_review": "still_required",
                "manifest": str(args.manifest),
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
